"""Versioned query loading, safe selection, validation, and portable exports.

Filter expressions are interpreted from a small allowlisted AST. They never execute
Python code. Python dataset plugins, in contrast, are explicitly trusted local code.
"""

from __future__ import annotations

import ast
import csv
import json
import random
from collections import Counter
from collections.abc import Iterable
from importlib.resources import files
from pathlib import Path
from typing import Any

from companybench.extensions import load_factory
from companybench.models import Query

SEGMENTS = {
    "simple": 'complexity == "L1"',
    "complex": 'complexity in ["L3", "L4"]',
    "technographic": '"technographic" in tags',
    "ai_companies": '"ai_companies" in tags',
    "ai_adoption": '"ai_adoption" in tags',
    "manufacturing": 'industry == "Manufacturing"',
    "signals": '"time_sensitive" in tags',
    "finite_reference": '"finite_reference" in tags',
}


def load_dataset(path: str | Path | None = None) -> list[Query]:
    """Load bundled queries, ``development``, JSON(L), or ``module:function``.

    Custom Python loaders execute trusted code. JSON and JSONL are data-only.
    Query indexes remain global identifiers and are never renumbered on selection.
    """
    if path is None or str(path) in {"default", "catalogue", "development"}:
        name = "development.jsonl" if str(path) == "development" else "queries.jsonl"
        text = files("companybench").joinpath("data", name).read_text(encoding="utf-8")
        raw: Any = [json.loads(line) for line in text.splitlines() if line.strip()]
    elif ":" in str(path) and not Path(path).exists():
        raw = list(load_factory(str(path))())
    else:
        dataset_path = Path(path)
        text = dataset_path.read_text(encoding="utf-8")
        if dataset_path.suffix.lower() == ".jsonl":
            raw = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            raw = json.loads(text)
            if isinstance(raw, dict):
                raw = raw.get("queries")
    if not isinstance(raw, list):
        raise ValueError("A dataset must contain a list of query objects")
    queries = [item if isinstance(item, Query) else Query.model_validate(item) for item in raw]
    errors = validate_dataset(queries)
    if errors:
        raise ValueError("Invalid dataset: " + "; ".join(errors))
    return sorted(queries, key=lambda query: query.index)


def validate_dataset(queries: Iterable[Query]) -> list[str]:
    """Return actionable semantic errors; allow noncontiguous custom indexes."""
    records = list(queries)
    errors = []
    for field in ("id", "index"):
        counts = Counter(getattr(query, field) for query in records)
        for value, count in counts.items():
            if count > 1:
                errors.append(f"Duplicate query {field}: {value}")
    for query in records:
        if not query.query.strip() or not query.acceptance.strip():
            errors.append(f"{query.id}: query and acceptance must be nonempty")
        if query.complexity not in {"L1", "L2", "L3", "L4"}:
            errors.append(f"{query.id}: complexity must be L1, L2, L3, or L4")
        if any(not condition.description.strip() for condition in query.conditions):
            errors.append(f"{query.id}: conditions must be nonempty")
        if getattr(query, "track", "public_web") != "public_web":
            errors.append(f"{query.id}: this release supports public_web queries only")
        if query.reference and query.reference.exhaustive:
            ref = query.reference
            if not ref.source_urls or not ref.as_of or not ref.notes:
                errors.append(
                    f"{query.id}: exhaustive reference requires source URLs, as_of, and scope notes"
                )
            names = [company.name.casefold() for company in ref.companies]
            if len(names) != len(set(names)):
                errors.append(f"{query.id}: duplicate reference company names")
    return errors


def _interpret(node: ast.AST, fields: dict[str, Any]) -> Any:
    if isinstance(node, ast.Expression):
        return _interpret(node.body, fields)
    if isinstance(node, ast.Name):
        if node.id not in fields:
            raise ValueError(f"Unknown filter field: {node.id}")
        return fields[node.id]
    if isinstance(node, ast.Constant) and isinstance(
        node.value, (str, int, float, bool, type(None))
    ):
        return node.value
    if isinstance(node, (ast.List, ast.Tuple, ast.Set)):
        return [_interpret(item, fields) for item in node.elts]
    if isinstance(node, ast.BoolOp) and isinstance(node.op, (ast.And, ast.Or)):
        values = [_interpret(item, fields) for item in node.values]
        return all(values) if isinstance(node.op, ast.And) else any(values)
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.Not):
        return not _interpret(node.operand, fields)
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
        value = _interpret(node.operand, fields)
        if type(value) not in {int, float}:
            raise ValueError("Signed filter literals must be numbers")
        return -value if isinstance(node.op, ast.USub) else value
    if isinstance(node, ast.Compare):
        left = _interpret(node.left, fields)
        matches = []
        for op, operand in zip(node.ops, node.comparators, strict=True):
            right = _interpret(operand, fields)
            if isinstance(op, ast.Eq):
                result = left == right
            elif isinstance(op, ast.NotEq):
                result = left != right
            elif isinstance(op, ast.In):
                result = left in right
            elif isinstance(op, ast.NotIn):
                result = left not in right
            elif isinstance(op, ast.Lt):
                result = left < right
            elif isinstance(op, ast.LtE):
                result = left <= right
            elif isinstance(op, ast.Gt):
                result = left > right
            elif isinstance(op, ast.GtE):
                result = left >= right
            else:
                raise ValueError("Unsupported filter comparison")
            matches.append(result)
            left = right
        return all(matches)
    raise ValueError(f"Unsupported filter syntax: {type(node).__name__}")


def select_queries(
    queries: Iterable[Query],
    ids: Iterable[str] | None = None,
    indices: Iterable[int] | None = None,
    index_range: str | tuple[int, int] | None = None,
    where: str | Iterable[str] | None = None,
    segments: str | Iterable[str] | None = None,
    sample: int | None = None,
    seed: int = 42,
) -> list[Query]:
    """Union explicit selectors, intersect filters, then sample without replacement."""
    records = sorted(queries, key=lambda query: query.index)
    field_names = set(Query.model_fields)
    field_names.update(key for query in records for key in query.model_dump())
    known_ids = {query.id for query in records}
    known_indices = {query.index for query in records}
    chosen_ids = [ids] if isinstance(ids, str) else list(ids or [])
    chosen_indices = list(indices) if indices is not None else []
    if len(chosen_ids) != len(set(chosen_ids)) or len(chosen_indices) != len(set(chosen_indices)):
        raise ValueError("Duplicate query selectors are not allowed")
    if set(chosen_ids) - known_ids:
        raise ValueError(f"Unknown query IDs: {sorted(set(chosen_ids) - known_ids)}")
    if any(type(index) is not int or index not in known_indices for index in chosen_indices):
        raise ValueError("Query indices must be existing one-based integer indexes")
    has_selector = ids is not None or indices is not None or index_range is not None
    if index_range is not None:
        try:
            start, end = (
                (int(value) for value in index_range.split(":"))
                if isinstance(index_range, str)
                else index_range
            )
        except (TypeError, ValueError) as error:
            raise ValueError("Range must be inclusive START:END") from error
        if (
            type(start) is not int
            or type(end) is not int
            or start > end
            or start not in known_indices
            or end not in known_indices
        ):
            raise ValueError("Range must have existing one-based endpoints with START <= END")
        chosen_indices.extend(range(start, end + 1))
    if has_selector:
        selected_indices = set(chosen_indices)
        records = [
            query for query in records if query.id in chosen_ids or query.index in selected_indices
        ]
    expressions = [where] if isinstance(where, str) else list(where or [])
    segment_names = [segments] if isinstance(segments, str) else list(segments or [])
    for name in segment_names:
        if name not in SEGMENTS:
            raise ValueError(f"Unknown segment {name!r}; choose from {', '.join(SEGMENTS)}")
        expressions.append(SEGMENTS[name])
    for expression in expressions:
        if len(expression) > 4096:
            raise ValueError("Filter exceeds the 4096-character limit")
        try:
            parsed = ast.parse(expression, mode="eval")
            if len(list(ast.walk(parsed))) > 200:
                raise ValueError("Filter is too complex")
            allowed = (
                ast.Expression,
                ast.Name,
                ast.Load,
                ast.Constant,
                ast.List,
                ast.Tuple,
                ast.Set,
                ast.BoolOp,
                ast.And,
                ast.Or,
                ast.UnaryOp,
                ast.Not,
                ast.USub,
                ast.UAdd,
                ast.Compare,
                ast.Eq,
                ast.NotEq,
                ast.In,
                ast.NotIn,
                ast.Lt,
                ast.LtE,
                ast.Gt,
                ast.GtE,
            )
            for node in ast.walk(parsed):
                if not isinstance(node, allowed):
                    raise ValueError(f"Unsupported filter syntax: {type(node).__name__}")
                if isinstance(node, ast.Name) and node.id not in field_names:
                    raise ValueError(f"Unknown filter field: {node.id}")
            records = [
                query
                for query in records
                if _interpret(parsed, dict.fromkeys(field_names) | query.model_dump())
            ]
        except (SyntaxError, TypeError, RecursionError) as error:
            raise ValueError(f"Invalid filter {expression!r}: {error}") from error
    if sample is not None:
        if type(sample) is not int or sample < 0 or sample > len(records):
            raise ValueError(f"Sample must be between 0 and {len(records)} after filtering")
        records = sorted(random.Random(seed).sample(records, sample), key=lambda query: query.index)
    return records


def export_dataset(queries: Iterable[Query], path: str | Path) -> Path:
    """Export JSONL/JSON, flat CSV, or an optional formatted XLSX catalogue."""
    destination = Path(path)
    records = [query.model_dump(mode="json") for query in queries]
    extension = destination.suffix.lower()
    if extension not in {".jsonl", ".json", ".csv", ".xlsx"}:
        raise ValueError("Export extension must be .jsonl, .json, .csv, or .xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    if extension == ".json":
        destination.write_text(
            json.dumps({"queries": records}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    elif extension == ".jsonl":
        destination.write_text(
            "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records),
            encoding="utf-8",
        )
    else:
        columns = list(dict.fromkeys(key for record in records for key in record))
        flat = [
            {
                key: json.dumps(value, ensure_ascii=False)
                if isinstance(value, (list, dict))
                else value
                for key, value in record.items()
            }
            for record in records
        ]
        if extension == ".csv":
            with destination.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=columns, lineterminator="\n")
                writer.writerow(
                    {
                        key: "'" + key if key.startswith(("=", "+", "-", "@")) else key
                        for key in columns
                    }
                )
                writer.writerows(
                    {
                        key: "'" + value
                        if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
                        else value
                        for key, value in row.items()
                    }
                    for row in flat
                )
        else:
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment, Font, PatternFill
            except ImportError as error:
                raise RuntimeError(
                    "XLSX export requires pip install 'companybench[excel]'"
                ) from error
            workbook = Workbook()
            sheet = workbook.active
            assert sheet is not None
            sheet.title = "Catalogue"
            sheet.append(columns)
            for row in flat:
                sheet.append([row.get(column) for column in columns])
            sheet.freeze_panes = "D2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.data_type = "s"
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="12332B")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column in sheet.columns:
                header = column[0]
                sheet.column_dimensions[header.column_letter].width = (
                    55 if header.value in {"query", "acceptance", "conditions"} else 24
                )
            workbook.save(destination)
    return destination
