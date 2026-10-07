import json

import pytest

from companybench.datasets import export_dataset, load_dataset, select_queries, validate_dataset


def test_packaged_catalogue_and_references():
    queries = load_dataset()
    assert len(queries) == 250
    assert [q.index for q in queries] == list(range(1, 251))
    assert not {f"CSB-{n:03}" for n in range(96, 121)} & {q.id for q in queries}
    assert len([q for q in queries if q.reference and q.reference.exhaustive]) == 10
    assert validate_dataset(queries) == []
    assert len(load_dataset("development")) == 12
    assert all(q.id.startswith("DEV-") for q in load_dataset("development"))


def test_selection_order_and_inclusive_global_indices():
    queries = load_dataset()
    assert [q.index for q in select_queries(queries, index_range="100:150")] == list(
        range(100, 151)
    )
    assert [q.index for q in select_queries(queries, indices=[142, 12, 101])] == [12, 101, 142]
    assert [
        q.id
        for q in select_queries(queries, ids=["CSB-001", "CSB-251"], where='complexity == "L1"')
    ] == ["CSB-001", "CSB-251"]
    a = select_queries(queries, segments=["simple"], sample=5, seed=13)
    assert a == select_queries(queries, segments=["simple"], sample=5, seed=13)
    assert all(q.complexity == "L1" for q in a)


def test_restricted_filters_preserve_boolean_and_membership_semantics():
    queries = load_dataset()
    selected = select_queries(
        queries, where=['industry == "Manufacturing"', 'complexity in ["L1", "L2"]']
    )
    assert selected and all(
        q.industry == "Manufacturing" and q.complexity in {"L1", "L2"} for q in selected
    )
    selected = select_queries(queries, where='"ai_companies" in tags or "ai_adoption" in tags')
    assert selected
    assert all("ai_companies" in q.tags or "ai_adoption" in q.tags for q in selected)
    for expression in [
        '__import__("os").system("true")',
        "query.__class__",
        "[q for q in tags]",
        "unknown_field == 1",
    ]:
        with pytest.raises(ValueError):
            select_queries(queries, where=expression)


def test_filter_on_generator_and_empty_result():
    queries = load_dataset()
    assert select_queries(iter(queries), where=["index < 0", '"ai_companies" in tags']) == []
    assert select_queries([], where='complexity == "L1"') == []
    with pytest.raises(ValueError, match="Unsupported"):
        select_queries([], where='__import__("os")')


def test_retained_semantic_condition_preserves_alternatives():
    queries = {query.id: query for query in load_dataset()}
    query = queries["CSB-034"]
    assert "United States or Canada" in query.conditions[0].description
    assert query.query in query.conditions[0].description
    assert query.acceptance in query.conditions[0].description
    assert queries["CSB-121"].verification_methods == ["public_html", "public_text"]
    assert "geospatial" in queries["CSB-198"].verification_methods


@pytest.mark.parametrize(
    "options",
    [
        {"indices": [0]},
        {"indices": [999]},
        {"indices": [True]},
        {"ids": ["MISSING"]},
        {"index_range": "150:100"},
        {"sample": -1},
        {"sample": 999},
        {"segments": ["missing"]},
        {"indices": [1, 1]},
    ],
)
def test_invalid_selection_is_not_silently_ignored(options):
    with pytest.raises(ValueError):
        select_queries(load_dataset(), **options)


def test_json_and_jsonl_round_trip(tmp_path):
    queries = load_dataset()[:3]
    for extension in ["json", "jsonl"]:
        path = tmp_path / f"queries.{extension}"
        export_dataset(queries, path)
        assert load_dataset(path) == queries
    data = tmp_path / "invalid.json"
    data.write_text(json.dumps([queries[0].model_dump(), queries[0].model_dump()]))
    with pytest.raises(ValueError, match="Duplicate"):
        load_dataset(data)


def test_reference_requires_traceable_exhaustive_evidence():
    query = next(q for q in load_dataset() if q.reference)
    bad = query.model_copy(deep=True)
    bad.reference.source_urls = []
    assert any("source" in error.lower() for error in validate_dataset([bad]))


def test_python_contribution_loader(tmp_path):
    module = tmp_path / "queries.py"
    module.write_text(
        "from companybench.datasets import load_dataset\ndef cases():\n    return load_dataset()[:2]\n"
    )
    assert len(load_dataset(f"{module}:cases")) == 2


def test_csv_and_xlsx_export(tmp_path):
    import csv

    from openpyxl import load_workbook

    queries = load_dataset()[:2]
    export_dataset(queries, tmp_path / "catalogue.csv")
    with (tmp_path / "catalogue.csv").open() as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 2
    assert json.loads(rows[0]["conditions"])
    export_dataset(queries, tmp_path / "catalogue.xlsx")
    workbook = load_workbook(tmp_path / "catalogue.xlsx")
    assert workbook["Catalogue"].max_row == 3
